import pytest
import os
import csv
import openpyxl
from datetime import datetime
import sys
import shutil # For cleanup_temp_files

# Adjust path to import from srz_kun.core
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from srz_kun.core.importer import (
    import_from_csv,
    import_from_excel,
    ImportFormatError,
    EXPECTED_CSV_EXCEL_HEADERS
)

# Define paths to sample files
# Assumes pytest is run from the project root where 'examples/' and 'tests/' are.
EXAMPLES_DIR = "examples"
SAMPLE_CSV_PATH = os.path.join(EXAMPLES_DIR, "sample_people.csv")
SAMPLE_EXCEL_PATH = os.path.join(EXAMPLES_DIR, "sample_people.xlsx")

# Base directory for temporary test files created by these tests
BASE_TEST_DIR = os.path.join(os.path.dirname(__file__), "temp_test_files_importer")


# Helper to create temporary test files
def create_temp_csv(filepath, headers, data_rows):
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(data_rows)

def create_temp_excel(filepath, headers, data_rows):
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in data_rows:
        processed_row = []
        for item_idx, item_val in enumerate(row):
            # For Excel, dates are often better as datetime objects or ISO strings
            # if headers[item_idx] in ["生年月日", "死亡年月日"] and isinstance(item_val, str) and item_val:
            #     try:
            #         # Attempt to parse common string formats if needed, but importer handles this
            #         # For test creation, direct datetime or ISO string is usually sufficient
            #         dt_obj = datetime.strptime(item_val, "%Y-%m-%d")
            #         processed_row.append(dt_obj)
            #         continue
            #     except ValueError:
            #         pass # Fallback to append as is
            processed_row.append(item_val)
        ws.append(processed_row)
    wb.save(filepath)

# --- Test Successful Imports ---
def test_import_from_csv_success(request):
    """Test successful import from the sample CSV file."""
    # Create the sample CSV if it doesn't exist (e.g. in a clean CI environment)
    if not os.path.exists(SAMPLE_CSV_PATH):
        print(f"Warning: {SAMPLE_CSV_PATH} not found, creating dummy for test.")
        sample_headers = ["ID", "氏名", "続柄", "生年月日", "死亡年月日", "本籍地", "住所", "生死", "相続放棄"]
        sample_data = [
            ["", "山田 太郎", "被相続人", "1950-01-01", "2023-10-26", "東京都千代田区", "東京都新宿区", "死亡", "なし"],
            ["P02", "山田 花子", "妻", "1955/05/10", "", "東京都千代田区", "東京都新宿区", "生存", "なし"],
            ["", "山田 一郎", "長男", "1980.08.15", "", "", "東京都中野区", "生存", "あり"],
            ["", "佐藤 良子", "長女", "1982-11-20", "", "神奈川県横浜市", "", "生存", "なし"],
            ["", "田中 次郎", "父", "1925-03-03", "2000-01-15", "長野県松本市", "長野県松本市", "死亡", "なし"]
        ]
        create_temp_csv(SAMPLE_CSV_PATH, sample_headers,sample_data)

    assert os.path.exists(SAMPLE_CSV_PATH), f"Sample CSV file {SAMPLE_CSV_PATH} still not found."

    data = import_from_csv(SAMPLE_CSV_PATH)
    assert isinstance(data, list)
    assert len(data) == 5

    person1 = data[0]
    assert person1["name"] == "山田 太郎"
    assert person1["relationship_to_deceased"] == "被相続人"
    assert person1["date_of_birth"] == "1950-01-01"
    assert person1["date_of_death"] == "2023-10-26"
    assert person1["is_alive"] is False
    assert person1["waived_inheritance"] is False

    person2 = data[1]
    assert person2["name"] == "山田 花子"
    assert person2["id"] == "P02"
    assert person2["date_of_birth"] == "1955-05-10"
    assert person2["is_alive"] is True

    person3 = data[2]
    assert person3["name"] == "山田 一郎"
    assert person3["date_of_birth"] == "1980-08-15"
    assert person3["waived_inheritance"] is True


def test_import_from_excel_success(request):
    """Test successful import from the sample Excel file."""
    if not os.path.exists(SAMPLE_EXCEL_PATH):
        print(f"Warning: {SAMPLE_EXCEL_PATH} not found, creating dummy for test.")
        sample_headers = ["ID", "氏名", "続柄", "生年月日", "死亡年月日", "本籍地", "住所", "生死", "相続放棄"]
        sample_data = [
            ["", "山田 太郎", "被相続人", datetime(1950, 1, 1), datetime(2023, 10, 26), "東京都千代田区", "東京都新宿区", "死亡", "なし"],
            ["P02", "山田 花子", "妻", datetime(1955, 5, 10), None, "東京都千代田区", "東京都新宿区", "生存", "なし"],
            ["", "山田 一郎", "長男", datetime(1980, 8, 15), None, "", "東京都中野区", "生存", "あり"],
            ["", "佐藤 良子", "長女", datetime(1982, 11, 20), None, "神奈川県横浜市", "", "生存", "なし"],
            ["", "田中 次郎", "父", datetime(1925, 3, 3), datetime(2000, 1, 15), "長野県松本市", "長野県松本市", "死亡", "なし"]
        ]
        create_temp_excel(SAMPLE_EXCEL_PATH, sample_headers, sample_data)

    assert os.path.exists(SAMPLE_EXCEL_PATH), f"Sample Excel file {SAMPLE_EXCEL_PATH} still not found."

    data = import_from_excel(SAMPLE_EXCEL_PATH)
    assert isinstance(data, list)
    assert len(data) == 5

    person1 = data[0]
    assert person1["name"] == "山田 太郎"
    assert person1["relationship_to_deceased"] == "被相続人"
    assert person1["date_of_birth"] == "1950-01-01"
    assert person1["date_of_death"] == "2023-10-26"
    assert person1["is_alive"] is False

    person2 = data[1]
    assert person2["id"] == "P02"
    assert person2["date_of_birth"] == "1955-05-10"
    assert person2["is_alive"] is True


# --- Test Error Handling ---
@pytest.fixture(scope="session", autouse=True)
def manage_temp_test_dir(request):
    """Creates and cleans up temporary test file directory for the session."""
    if os.path.exists(BASE_TEST_DIR): # Cleanup from previous run if any
        shutil.rmtree(BASE_TEST_DIR)
    os.makedirs(BASE_TEST_DIR, exist_ok=True)

    def remove_temp_dir():
        if os.path.exists(BASE_TEST_DIR):
            shutil.rmtree(BASE_TEST_DIR)
            # print(f"Cleaned up {BASE_TEST_DIR}")
    request.addfinalizer(remove_temp_dir)


def test_import_csv_file_not_found():
    with pytest.raises(FileNotFoundError):
        import_from_csv(os.path.join(BASE_TEST_DIR, "non_existent_file.csv"))

def test_import_excel_file_not_found():
    with pytest.raises(FileNotFoundError):
        import_from_excel(os.path.join(BASE_TEST_DIR, "non_existent_file.xlsx"))

def test_import_csv_missing_required_header():
    headers = ["氏名", "生年月日"] # Missing "続柄"
    data_rows = [["テスト 太郎", "1990-01-01"]]
    temp_file = os.path.join(BASE_TEST_DIR, "missing_header.csv")
    create_temp_csv(temp_file, headers, data_rows)
    with pytest.raises(ImportFormatError, match="必須項目「続柄」が見つかりません"):
        import_from_csv(temp_file)

def test_import_excel_missing_required_header():
    headers = ["氏名", "生年月日"] # Missing "続柄"
    data_rows = [["テスト 花子", datetime(1995,2,2)]] # Excel dates are often datetime
    temp_file = os.path.join(BASE_TEST_DIR, "missing_header.xlsx")
    create_temp_excel(temp_file, headers, data_rows)
    with pytest.raises(ImportFormatError, match="必須項目「続柄」が見つかりません"):
        import_from_excel(temp_file)

def test_import_csv_invalid_is_alive_value():
    headers = list(EXPECTED_CSV_EXCEL_HEADERS.keys()) # Use all expected headers
    data_rows = [["", "テスト 次郎", "子", "2000-01-01", "", "", "", "不明", "なし"]]
    temp_file = os.path.join(BASE_TEST_DIR, "invalid_is_alive.csv")
    create_temp_csv(temp_file, headers, data_rows)
    with pytest.raises(ImportFormatError, match="「生死」の値が不正です: 「不明」"):
        import_from_csv(temp_file)

def test_import_csv_invalid_date_format():
    headers = ["氏名", "続柄", "生年月日"]
    data_rows = [["テスト 三郎", "孫", "01-01-1990"]]
    temp_file = os.path.join(BASE_TEST_DIR, "invalid_date.csv")
    create_temp_csv(temp_file, headers, data_rows)
    with pytest.raises(ImportFormatError, match="日付フィールド「生年月日」の値「01-01-1990」は YYYY-MM-DD .* 形式で入力してください。"):
        import_from_csv(temp_file)

def test_import_excel_invalid_waived_value():
    headers = list(EXPECTED_CSV_EXCEL_HEADERS.keys())
    data_rows = [["", "テスト 四郎", "兄弟", "2000-01-01", "", "", "", "生存", "たぶん"]]
    temp_file = os.path.join(BASE_TEST_DIR, "invalid_waived.xlsx")
    create_temp_excel(temp_file, headers, data_rows)
    with pytest.raises(ImportFormatError, match="「相続放棄」の値が不正です: 「たぶん」"):
        import_from_excel(temp_file)

def test_import_excel_non_excel_file():
    temp_file = os.path.join(BASE_TEST_DIR, "not_an_excel.xlsx")
    with open(temp_file, "w") as f:
        f.write("This is not an excel file.")
    with pytest.raises(ImportFormatError, match="Excelファイル形式が無効です"):
        import_from_excel(temp_file)

def test_empty_csv_file():
    temp_file = os.path.join(BASE_TEST_DIR, "empty.csv")
    create_temp_csv(temp_file, [], [])
    with pytest.raises(ImportFormatError, match="CSVファイルにヘッダー行が見つかりません。"):
        import_from_csv(temp_file)

def test_empty_excel_file():
    temp_file = os.path.join(BASE_TEST_DIR, "empty.xlsx")
    wb = openpyxl.Workbook()
    # Save an empty workbook (or one with just an empty sheet)
    wb.save(temp_file)
    # The current check is max_row < 1, which an empty sheet might pass.
    # Let's make it more specific: no headers.
    # Or, if it has headers but no data, it should return empty list not error.
    # Test case: header but no data rows
    create_temp_excel(temp_file, list(EXPECTED_CSV_EXCEL_HEADERS.keys()), [])
    data = import_from_excel(temp_file)
    assert data == []

    # Test case: truly empty sheet (no headers)
    wb_empty = openpyxl.Workbook()
    wb_empty.save(temp_file)
    with pytest.raises(ImportFormatError, match="Excelファイルからヘッダー行を読み取れませんでした。"):
         import_from_excel(temp_file)

def test_csv_date_normalization_slashes():
    headers = ["氏名", "続柄", "生年月日"]
    data_rows = [["スラッシュ 日付", "子", "1999/12/31"]]
    temp_file = os.path.join(BASE_TEST_DIR, "date_slash.csv")
    create_temp_csv(temp_file, headers, data_rows)
    data = import_from_csv(temp_file)
    assert data[0]["date_of_birth"] == "1999-12-31"

def test_csv_date_normalization_dots():
    headers = ["氏名", "続柄", "生年月日"]
    data_rows = [["ドット 日付", "孫", "1998.01.05"]]
    temp_file = os.path.join(BASE_TEST_DIR, "date_dot.csv")
    create_temp_csv(temp_file, headers, data_rows)
    data = import_from_csv(temp_file)
    assert data[0]["date_of_birth"] == "1998-01-05"

```
